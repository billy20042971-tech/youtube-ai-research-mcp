import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from textwrap import wrap

try:
    from scripts.daily_summary import build_summary, report_date_taipei
    from scripts.daily_summary_evidence import evidence_metadata
except ModuleNotFoundError:
    # Support direct execution as: python scripts/build_daily_reports.py
    from daily_summary import build_summary, report_date_taipei
    from daily_summary_evidence import evidence_metadata


def safe_int(value):
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


src = json.loads(Path("data/latest.json").read_text(encoding="utf-8"))
recent = src.get("recent_72h_sample", {})
enriched_by_id = {
    v.get("video_id"): v for v in recent.get("videos", [])
    if isinstance(v, dict) and v.get("video_id")
}
videos = []
for query, data in recent.get("per_query", {}).items():
    for raw in data.get("videos", []):
        if not isinstance(raw, dict):
            continue
        video = dict(raw)
        video["query"] = query
        enriched = enriched_by_id.get(video.get("video_id"), {})
        for key in ("title", "channel_id", "channel_title", "published_at", "description",
                    "tags", "views", "likes", "comments", "url", "comments_obtained_count"):
            if key in enriched:
                video[key] = enriched[key]
        comments = video.get("comments")
        if not isinstance(comments, list):
            video["comments_count"] = safe_int(comments)
            video["comments"] = []
        else:
            video["comments_count"] = len(comments)
        videos.append(video)

raw_video_count = sum(1 for video in videos if video.get("video_id"))
deduped = {}
for video in videos:
    if video.get("video_id"):
        deduped[video["video_id"]] = video
videos = sorted(deduped.values(), key=lambda v: safe_int(v.get("views")), reverse=True)
now = datetime.now(timezone.utc)
date = report_date_taipei(now)
out = Path("reports") / date
out.mkdir(parents=True, exist_ok=True)
summary = build_summary(videos, generated_at=src.get("generated_at"))
evidence = evidence_metadata(videos, src.get("generated_at"))
evidence["duplicate_count"] = max(0, raw_video_count - len(videos))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
available_fonts = {f.name for f in font_manager.fontManager.ttflist}
cjk_candidates = ["Noto Sans CJK TC", "Noto Sans CJK HK", "Noto Sans CJK SC", "Noto Sans CJK JP"]
chosen_font = next((name for name in cjk_candidates if name in available_fonts), None)
if not chosen_font:
    raise RuntimeError("No supported CJK font found; install fonts-noto-cjk.")
plt.rcParams["font.family"] = chosen_font
plt.rcParams["axes.unicode_minus"] = False

top = videos[:10]
def chart_label(title):
    text = re.sub(r"\s+", " ", str(title or "")).strip() or "(無標題)"
    lines = wrap(text, width=15, break_long_words=True, break_on_hyphens=False)
    if len(lines) > 2:
        lines = [lines[0], lines[1][:13] + "…"]
    return "\n".join(lines)

labels = [chart_label(v.get("title")) for v in top][::-1]
values = [safe_int(v.get("views")) for v in top][::-1]
fig_height = max(7.0, 0.82 * len(labels) + 1.8)
fig, ax = plt.subplots(figsize=(14, fig_height))
bars = ax.barh(labels, values)
ax.set_title("YouTube AI 研究雷達｜觀看數前 10 名（僅為曝光訊號）", fontsize=15, pad=16)
ax.set_xlabel("觀看次數")
ax.tick_params(axis="y", labelsize=9)
ax.tick_params(axis="x", labelsize=9)
ax.grid(axis="x", linestyle=":", alpha=0.35)
ax.set_axisbelow(True)
max_value = max(values, default=0)
ax.set_xlim(0, max_value * 1.18 if max_value else 1)
for bar, value in zip(bars, values):
    ax.annotate(
        f"{value:,}",
        xy=(bar.get_width(), bar.get_y() + bar.get_height() / 2),
        xytext=(5, 0),
        textcoords="offset points",
        ha="left",
        va="center",
        fontsize=8,
        clip_on=False,
    )
fig.subplots_adjust(left=0.43, right=0.96, top=0.94, bottom=0.08)
chart_path = out / "top10_views.png"
fig.savefig(chart_path, dpi=180, bbox_inches="tight", facecolor="white")
plt.close(fig)

from PIL import Image
with Image.open(chart_path) as image:
    image.verify()
with Image.open(chart_path) as image:
    if image.width < 1000 or image.height < 600:
        raise RuntimeError(f"Chart dimensions too small: {image.size}")
if chart_path.stat().st_size < 20000:
    raise RuntimeError(f"Chart file suspiciously small: {chart_path.stat().st_size} bytes")

channels = Counter(v.get("channel_title", "") for v in videos if v.get("channel_title"))
queries = Counter(v.get("query", "") for v in videos if v.get("query"))

from docx import Document
from docx.shared import Inches
doc = Document()
doc.add_heading(f"YouTube AI 趨勢研究雷達｜{date}", 0)
doc.add_paragraph(f"產出時間：{now.isoformat()}｜近 72 小時去重樣本：{len(videos)}")
doc.add_heading("1. 一分鐘決策摘要", level=1)
doc.add_paragraph(summary["headline"])
for i, signal in enumerate(summary["signals"], 1):
    doc.add_heading(f"{i}. {signal['topic']}｜{signal['classification']}", level=2)
    doc.add_paragraph(
        f"樣本 {signal['video_count']} 部／頻道 {signal['channel_count']} 個；"
        f"證據強度：{signal['confidence']}。{signal['evidence']}。"
    )
    doc.add_paragraph(
        f"代表影片：{signal['representative_title']}｜觀看 {signal['views']:,}｜"
        f"發布：{signal['published_at']}\n{signal['representative_url']}"
    )
doc.add_paragraph("行動建議：目前資料只支持「持續觀察」；不要僅憑觀看數採購工具、改變教學或宣稱投資／收入效果。")
doc.add_heading("證據限制", level=2)
for limitation in summary["limitations"]:
    doc.add_paragraph("• " + limitation)
doc.add_paragraph(
    f"證據狀態：{evidence['trend_status']}｜去重樣本 {evidence['sample_count']} 部｜"
    f"跨查詢重複 {evidence['duplicate_count']} 部｜超過72h {evidence['stale_count']} 部｜"
    f"日期不可判定 {evidence['freshness_unknown_count']} 部。"
)
for signal in evidence["signals"]:
    doc.add_paragraph(
        f"證據卡：{signal['title']}｜發布：{signal['published_at']}｜新鮮度：{signal['freshness']}｜"
        f"來源獨立性：{signal['source_independence']}｜工作關聯：{'、'.join(signal['work_relevance'])}｜"
        f"建議：{signal['action']}｜{signal['url']}"
    )
for limitation in evidence["limitations"]:
    doc.add_paragraph("證據限制：" + limitation)
if summary["deep_dive"]:
    doc.add_heading("最多 1 部候選深讀影片", level=2)
    doc.add_paragraph(
        f"{summary['deep_dive']['title']}\n{summary['deep_dive']['url']}\n"
        f"{summary['deep_dive']['reason']} 時間戳：未提供可靠時間戳，不臆測。"
    )

doc.add_heading("2. 觀看數前 10 名（曝光訊號，不等於趨勢證明）", level=1)
for i, video in enumerate(top, 1):
    doc.add_paragraph(
        f"{i}. {video.get('title') or '(無標題)'}｜{video.get('channel_title') or '頻道未提供'}｜"
        f"{safe_int(video.get('views')):,} views｜{video.get('published_at') or '發布時間未提供'}\n"
        f"{video.get('url') or ''}"
    )
doc.add_picture(str(chart_path), width=Inches(6.4))

doc.add_heading("3. 留言抽樣（非代表性民調）", level=1)
for video in top[:8]:
    comments = video.get("comments") if isinstance(video.get("comments"), list) else []
    doc.add_paragraph(f"{video.get('title') or '(無標題)'}｜實際取得留言 {safe_int(video.get('comments_obtained_count', len(comments)))} 則")
    for comment in comments[:3]:
        text = comment.get("text", "") if isinstance(comment, dict) else str(comment)
        if text:
            doc.add_paragraph("・" + str(text).replace("\n", " ")[:220])

doc.add_heading("4. AI OS 實作觀察（通用框架，非本次樣本的趨勢結論）", level=1)
for item in [
    "Agent：規劃、工具選擇、執行與停止條件；高風險動作需人工授權。",
    "MCP：工具／資料連接層；外部文字視為不可信輸入。",
    "Verification：檢查資料新鮮度、輸出完整性與寫回讀回結果。",
    "State／Recovery：保存 execution_id、checkpoint 與失敗原因，避免重複副作用。",
    "Automation：排程是否準時與研究流程是否成功，必須分開驗收。",
]:
    doc.add_paragraph("• " + item)

word_path = out / f"YouTube_AI趨勢研究雷達_{date}_完整報告.docx"
doc.save(word_path)

from pptx import Presentation
from pptx.util import Inches as PInches
prs = Presentation()
def slide(title, body):
    page = prs.slides.add_slide(prs.slide_layouts[1])
    page.shapes.title.text = title
    page.placeholders[1].text = body

slide("YouTube AI 趨勢研究雷達", f"{date}｜近 72 小時去重樣本 {len(videos)} 部")
slide("一分鐘決策摘要", summary["headline"] + "\n" + "\n".join(
    f"{i}. {s['topic']}｜{s['confidence']}｜{s['video_count']} 部／{s['channel_count']} 頻道"
    for i, s in enumerate(summary["signals"], 1)
) + "\n" +
    f"證據：{evidence['trend_status']}｜去重樣本 {evidence['sample_count']}｜重複 {evidence['duplicate_count']}｜"
    f"超過72h {evidence['stale_count']}｜日期不可判定 {evidence['freshness_unknown_count']}\n" +
    "來源獨立性未驗證；工作關聯為關鍵字初篩。")
slide("證據限制", "\n".join("• " + x for x in summary["limitations"]))
slide("下一步", "目前採取：持續觀察。觀看數不是需求、收入、效果或可靠性的證明。\n有可驗證證據後，再決定是否安排實驗。")
slide("高觀看影片候選", "\n".join(
    f"{i}. {(v.get('title') or '(無標題)')[:55]}｜{safe_int(v.get('views')):,}"
    for i, v in enumerate(top[:7], 1)
))
page = prs.slides.add_slide(prs.slide_layouts[5])
page.shapes.add_picture(str(chart_path), PInches(0.6), PInches(1.1), width=PInches(8.8))
pptx_path = out / f"YouTube_AI趨勢研究雷達_{date}_簡報.pptx"
prs.save(pptx_path)

manifest = {
    "generated_at": now.isoformat(),
    "source": "YouTube Research MCP",
    "sample_count": len(videos),
    "summary_signal_count": len(summary["signals"]),
    "summary_status": "evidence_limited" if videos else "insufficient_sample",
    "evidence_summary": evidence,
    "word": str(word_path),
    "pptx": str(pptx_path),
    "chart": str(chart_path),
    "automation_status": "generated",
}
Path("reports/latest-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(manifest, ensure_ascii=False, indent=2))
